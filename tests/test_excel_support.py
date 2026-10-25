import importlib
from io import BytesIO
import json
import os
from pathlib import Path
from types import SimpleNamespace
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import Mock, patch
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
import pandas as pd
from streamlit.testing.v1 import AppTest

from excel_support import (
    excel_sheet_names, generate_excel_download, query_sales_total, read_excel_sheet,
    spreadsheet_answer, spreadsheet_schema, write_region_summary,
)


ROOT = Path(__file__).resolve().parents[1]


def employee_app():
    app = AppTest.from_file(str(ROOT / "employee_sql_chatbot.py"))
    app.session_state["data_source"] = "Employee database"
    return app


TEST_ENVIRONMENT = {
    key: value for key, value in os.environ.items()
    if not key.upper().startswith(("AZURE_", "SMTP_", "WORKFORCE_"))
}


def workbook_bytes(sheets=None):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in (sheets or {
        "Sales": [("Region", "Revenue", "Date"), ("West", 10, pd.Timestamp("2025-01-01")),
                  ("East", 20, pd.Timestamp("2025-02-01")), ("West", 30, None)],
        "Inventory": [("Product", "Stock"), ("Widget", 7)],
    }).items():
        worksheet = workbook.create_sheet(title)
        for row in rows:
            worksheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def plan(**changes):
    value = {"understood": True, "unhandled_conditions": [], "columns": [],
             "metrics": [{"function": "count", "column": "*"}], "filter": None}
    value.update(changes)
    return value


class ExcelFileTests(unittest.TestCase):
    def test_sheet_selection_and_typed_schema(self):
        content = workbook_bytes()
        self.assertEqual(excel_sheet_names(content), ["Sales", "Inventory"])
        sales = read_excel_sheet(content, "Sales")
        self.assertEqual(sales.shape, (3, 3))
        data, schema, names = spreadsheet_schema(sales)
        self.assertEqual(schema, {"c0": "text", "c1": "number", "c2": "date"})
        self.assertEqual(names["c1"], "Revenue")
        self.assertEqual(data.loc[0, "c2"], "2025-01-01")
        self.assertTrue(pd.isna(data.loc[2, "c2"]))
        self.assertEqual(read_excel_sheet(content, "Inventory").iloc[0].tolist(), ["Widget", 7])

    def test_invalid_archives_and_unknown_sheets_are_explicit(self):
        for content in (b"", b"not an excel file", b"x" * 10_000_001):
            with self.subTest(length=len(content)), self.assertRaises(ValueError):
                excel_sheet_names(content)
        with self.assertRaisesRegex(ValueError, "couldn't find the selected worksheet"):
            read_excel_sheet(workbook_bytes(), "Missing")
        expanded = BytesIO()
        with ZipFile(expanded, "w", compression=ZIP_DEFLATED) as archive:
            archive.writestr("oversized", b"x" * 50_000_001)
        with self.assertRaisesRegex(ValueError, "50 MB"):
            excel_sheet_names(expanded.getvalue())

    def test_invalid_headers_empty_sheets_and_extra_cells_are_rejected(self):
        for rows in (
            [], [("Header",)], [("Name", " name "), ("A", "B")],
            [(None, "Revenue"), ("West", 10)], [(42,), (1,)],
            [("Name",), ("A", "no header")], [tuple(f"c{i}" for i in range(101))],
        ):
            with self.subTest(rows=rows[:1]), self.assertRaises(ValueError):
                read_excel_sheet(workbook_bytes({"Sheet": rows}), "Sheet")

    def test_actual_row_limit(self):
        rows = [("Number",)] + [(value,) for value in range(50_001)]
        content = workbook_bytes({"Data": rows})
        with self.assertRaisesRegex(ValueError, "50,000"):
            read_excel_sheet(content, "Data")
        content = workbook_bytes({"Data": rows[:-1]})
        self.assertEqual(len(read_excel_sheet(content, "Data")), 50_000)
        content = workbook_bytes({"Data": [
            tuple(f"Column {index}" for index in range(100)), tuple(range(100)),
        ]})
        self.assertEqual(read_excel_sheet(content, "Data").shape, (1, 100))

    def test_trailing_empty_formatted_header_cells_do_not_invalidate_data(self):
        content = workbook_bytes({"products": [
            ("Product ID", "Product Name"), (1, "Widget"), (2, "Service"),
        ]})
        workbook = load_workbook(BytesIO(content))
        workbook["products"]["Z1"].font = Font(bold=True)
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        frame = read_excel_sheet(output.getvalue(), "products")
        self.assertEqual(frame.to_dict("list"), {"Product ID": [1, 2], "Product Name": ["Widget", "Service"]})

    def test_header_errors_identify_cells_without_renaming_or_discarding_data(self):
        for rows, cells in (
            ([(None, "Revenue"), ("North", 1)], ["A1"]),
            ([("ID", " id "), (1, 2)], ["A1", "B1"]),
            ([("ID",), (1, "unlabelled")], ["B2"]),
        ):
            with self.subTest(rows=rows):
                with self.assertRaises(ValueError) as caught:
                    read_excel_sheet(workbook_bytes({"products": rows}), "products")
                for cell in cells:
                    self.assertIn(cell, str(caught.exception))
                self.assertIn("products", str(caught.exception))

    def test_exports_preserve_numbers_dates_and_literal_formulas(self):
        frame = pd.DataFrame({
            "=Header": ["=1+1", "+cmd", "-cmd", "@SUM(A1)", "https://example.com", "00123"],
            "Number": [1.25, 2, 3, 4, 5, None],
        })
        content = generate_excel_download([("Bad/name", frame), ("bad:name", frame)])
        workbook = load_workbook(BytesIO(content), data_only=False)
        self.assertEqual(workbook.sheetnames, ["Bad_name", "bad_name (2)"])
        worksheet = workbook.worksheets[0]
        self.assertEqual(worksheet["A1"].value, "=Header")
        self.assertEqual(worksheet["A2"].value, "=1+1")
        self.assertEqual(worksheet["A2"].data_type, "s")
        self.assertEqual(worksheet["B2"].value, 1.25)
        self.assertEqual(worksheet["A7"].value, "00123")
        self.assertIsNone(worksheet["A6"].hyperlink)
        self.assertEqual(worksheet.freeze_panes, "A2")
        dated = read_excel_sheet(workbook_bytes(), "Sales")
        exported = generate_excel_download([("Dates", dated)])
        self.assertEqual(load_workbook(BytesIO(exported)).active["C2"].value.year, 2025)
        workbook.close()

    def test_full_size_export_preserves_last_row_and_timezone_dates(self):
        frame = pd.DataFrame({"ID": range(1, 50_001), "Label": ["=1+1"] * 50_000})
        content = generate_excel_download([("Complete results", frame)])
        workbook = load_workbook(BytesIO(content), read_only=True)
        rows = workbook.active.iter_rows(values_only=True)
        self.assertEqual(next(rows), ("ID", "Label"))
        count = 0
        last = None
        for last in rows:
            count += 1
        self.assertEqual(count, 50_000)
        self.assertEqual(last, (50_000, "=1+1"))
        workbook.close()
        dates = pd.DataFrame({"Date": pd.to_datetime(["2025-01-01T10:30:00Z", None])})
        workbook = load_workbook(BytesIO(generate_excel_download([("Dates", dates)])))
        self.assertEqual(workbook.active["A2"].value, pd.Timestamp("2025-01-01T10:30:00").to_pydatetime())
        self.assertIsNone(workbook.active["A3"].value)
        workbook.close()
    def test_sales_totals_variants_preserve_filters_and_missing_values(self):
        frame = pd.DataFrame({"Region": ["North", "North", "South"], "Revenue": [10.125, -2, None]})
        for question in (
            "What are the total sales for the North region?",
            "what is the total revenue in the NORTH region?",
            "Show me the sum of sales for the North region",
        ):
            with self.subTest(question=question):
                result = query_sales_total(question, frame)
                self.assertEqual(result.iloc[0, 0], 8.125)
                self.assertIn("8.125", spreadsheet_answer(result))
                self.assertIn("North region", spreadsheet_answer(result))
        missing = query_sales_total("Total sales for the South region", frame)
        self.assertTrue(pd.isna(missing.iloc[0, 0]))
        self.assertIn("No value is available", spreadsheet_answer(missing))
        for question in (
            "Total sales for the North region in 2025",
            "Total sales for the North and South region",
            "Total sales for the North region excluding refunds",
            "Total sales by region excluding refunds",
        ):
            with self.subTest(question=question):
                self.assertIsNone(query_sales_total(question, frame))
        text_sales = pd.DataFrame({"Region": ["North"], "Sales": ["100 dollars"]})
        with self.assertRaisesRegex(ValueError, "numeric values"):
            query_sales_total("Total sales for the North region", text_sales)

    def test_plain_answers_for_counts_scalar_grouped_and_empty_results(self):
        self.assertEqual(spreadsheet_answer(pd.DataFrame({"row_count": [3]})), "There are 3 rows.")
        self.assertEqual(spreadsheet_answer(pd.DataFrame({"row_count": [1]})), "There is 1 row.")
        self.assertEqual(spreadsheet_answer(pd.DataFrame({"sum Revenue": [1234.567]})),
                         "The total Revenue is 1,234.567.")
        grouped = spreadsheet_answer(pd.DataFrame({"Region": ["North", "South"], "sum Sales": [100, 200]}))
        self.assertIn("Region: North; total Sales: 100.", grouped)
        self.assertNotIn("sum Sales", grouped)
        self.assertIn("No matching rows", spreadsheet_answer(pd.DataFrame()))

    def test_region_summary_requires_explicit_confirmation_before_opening(self):
        with patch("excel_support.load_workbook") as load:
            for confirmed in (False, None, "yes", 1):
                with self.subTest(confirmed=confirmed), self.assertRaisesRegex(ValueError, "confirm"):
                    write_region_summary(workbook_bytes(), confirmed=confirmed)
            load.assert_not_called()

    def test_workbook_paths_load_without_modifying_the_original(self):
        content = workbook_bytes()
        with TemporaryDirectory() as directory:
            path = Path(directory) / "sales.xlsx"
            path.write_bytes(content)
            self.assertEqual(excel_sheet_names(path), ["Sales", "Inventory"])
            self.assertEqual(read_excel_sheet(str(path), "Sales").shape, (3, 3))
            backup, updated = write_region_summary(path, confirmed=True)
            self.assertEqual(path.read_bytes(), content)
            self.assertEqual(backup, content)
            self.assertIn("Region_Summary", excel_sheet_names(updated))

    def test_missing_invalid_and_directory_paths_have_friendly_messages(self):
        with TemporaryDirectory() as directory:
            missing = Path(directory) / "missing.xlsx"
            for operation in (
                lambda: excel_sheet_names(missing),
                lambda: read_excel_sheet(missing, "Sales"),
                lambda: write_region_summary(missing, confirmed=True),
            ):
                with self.subTest(operation=operation), self.assertRaisesRegex(ValueError, "Check the file path"):
                    operation()
            with self.assertRaisesRegex(ValueError, "points to a folder"):
                excel_sheet_names(Path(directory))
            with self.assertRaisesRegex(ValueError, "path isn't valid"):
                excel_sheet_names("invalid\x00path.xlsx")

    def test_locked_files_show_remedial_advice_without_exception_details(self):
        for operation in (
            lambda: excel_sheet_names(Path("locked.xlsx")),
            lambda: read_excel_sheet(Path("locked.xlsx"), "Sales"),
            lambda: write_region_summary(Path("locked.xlsx"), confirmed=True),
        ):
            with self.subTest(operation=operation), patch(
                "pathlib.Path.open", side_effect=PermissionError("PRIVATE_PATH_AND_TECHNICAL_DETAILS"),
            ):
                with self.assertRaisesRegex(ValueError, "may be locked") as caught:
                    operation()
                self.assertIn("Close it in Excel", str(caught.exception))
                self.assertNotIn("PRIVATE_PATH", str(caught.exception))
        error = OSError("PRIVATE_LOCK_DETAILS")
        error.winerror = 32
        with patch("pathlib.Path.open", side_effect=error), self.assertRaisesRegex(ValueError, "may be locked"):
            excel_sheet_names(Path("locked.xlsx"))

    def test_workbook_load_and_save_access_errors_are_friendly(self):
        content = workbook_bytes()
        with patch("excel_support.load_workbook", side_effect=PermissionError("PRIVATE_DETAILS")):
            with self.assertRaisesRegex(ValueError, "may be locked"):
                excel_sheet_names(content)
        with patch("openpyxl.workbook.workbook.Workbook.save", side_effect=PermissionError("PRIVATE_DETAILS")):
            with self.assertRaisesRegex(ValueError, "Close it in Excel"):
                write_region_summary(content, confirmed=True)
        self.assertEqual(excel_sheet_names(content), ["Sales", "Inventory"])

    def test_missing_sales_sheet_explains_how_to_fix_it(self):
        content = workbook_bytes({"Inventory": [("Product",), ("Widget",)]})
        for operation in (
            lambda: read_excel_sheet(content, "Sales"),
            lambda: write_region_summary(content, confirmed=True),
        ):
            with self.subTest(operation=operation), self.assertRaisesRegex(ValueError, "sheet named 'Sales'") as caught:
                operation()
            self.assertIn("save the file, and upload it again", str(caught.exception))

    def test_region_summary_backs_up_original_and_preserves_existing_sheets(self):
        content = workbook_bytes({
            "Sales": [("Region", "Sales"), ("North", 10.5), (" north ", -2), ("South", 20)],
            "Notes": [("Formula",), ("=SUM(Sales!B2:B4)",)],
        })
        original = load_workbook(BytesIO(content))
        original["Notes"]["A1"].font = Font(bold=True, color="FF0000")
        original["Notes"].sheet_state = "hidden"
        buffer = BytesIO()
        original.save(buffer)
        content = buffer.getvalue()
        backup, updated = write_region_summary(content, confirmed=True)
        self.assertEqual(backup, content)
        workbook = load_workbook(BytesIO(updated), data_only=False)
        self.assertEqual(workbook.sheetnames, ["Sales", "Notes", "Region_Summary"])
        self.assertEqual(list(workbook["Region_Summary"].values), [
            ("Region", "Total Sales"), ("North", 8.5), ("South", 20),
        ])
        self.assertEqual(list(workbook["Sales"].values), list(original["Sales"].values))
        self.assertEqual(workbook["Notes"]["A2"].value, "=SUM(Sales!B2:B4)")
        self.assertTrue(workbook["Notes"]["A1"].font.bold)
        self.assertEqual(workbook["Notes"]["A1"].font.color.rgb, "00FF0000")
        self.assertEqual(workbook["Notes"].sheet_state, "hidden")
        self.assertNotIn("Region_Summary", load_workbook(BytesIO(backup)).sheetnames)
        workbook.close()
        original.close()

    def test_region_summary_rejects_existing_sheet_and_invalid_schema(self):
        for name in ("Region_Summary", "region_summary"):
            content = workbook_bytes({
                "Sales": [("Region", "Sales"), ("North", 10)], name: [("Keep me",), (42,)],
            })
            with self.subTest(name=name), self.assertRaisesRegex(ValueError, "already exists"):
                write_region_summary(content, confirmed=True)
            self.assertEqual(load_workbook(BytesIO(content))[name]["A2"].value, 42)
        for sheets in (
            {"Inventory": [("Product",), ("Widget",)]},
            {"Sales": [("Sales",), (10,)]},
            {"Sales": [("Region", "Other"), ("North", 10)]},
            {"Sales": [("Region", "Sales"), ("North", "ten")]},
        ):
            with self.subTest(sheets=sheets), self.assertRaises(ValueError):
                write_region_summary(workbook_bytes(sheets), confirmed=True)

    def test_region_summary_uses_all_rows_missing_totals_and_literal_labels(self):
        content = workbook_bytes({"Sales": [
            ("Region", "Revenue"), *[("North", 1) for _ in range(101)],
            ("South", None), (None, 5), ("", 2), ("=1+1", 7),
        ]})
        original = load_workbook(BytesIO(content))
        original["Sales"].cell(row=original["Sales"].max_row, column=1).data_type = "s"
        buffer = BytesIO()
        original.save(buffer)
        content = buffer.getvalue()
        original.close()
        _, updated = write_region_summary(content, confirmed=True)
        workbook = load_workbook(BytesIO(updated))
        self.assertEqual(list(workbook["Region_Summary"].values), [
            ("Region", "Total Sales"), ("North", 101), ("South", None), (None, 7), ("=1+1", 7),
        ])
        self.assertEqual(workbook["Region_Summary"]["A5"].data_type, "s")
        workbook.close()


class ExcelChatTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        with patch("dotenv.load_dotenv"), patch.dict(os.environ, TEST_ENVIRONMENT, clear=True):
            cls.chatbot = importlib.import_module("employee_sql_chatbot")

    def setUp(self):
        self.enterContext(patch("dotenv.load_dotenv"))
        self.enterContext(patch.dict(os.environ, TEST_ENVIRONMENT, clear=True))
        self.enterContext(patch("socket.create_connection", side_effect=AssertionError("Unexpected network")))
        self.frame = read_excel_sheet(workbook_bytes(), "Sales")
        self.client = Mock()
        self.enterContext(patch.object(self.chatbot, "get_azure_openai_client", return_value=self.client))
        self.enterContext(patch.dict(os.environ, {"AZURE_OPENAI_DEPLOYMENT_MINI": "test"}))

    def ask(self, question, query_plan):
        self.client.chat.completions.create.return_value = SimpleNamespace(choices=[
            SimpleNamespace(finish_reason="stop", message=SimpleNamespace(content=json.dumps(query_plan)))
        ])
        return self.chatbot.query_excel_sheet(question, self.frame)

    def test_count_and_preview_do_not_call_azure_or_employee_database(self):
        with patch.object(self.chatbot, "get_engine", side_effect=AssertionError("Employee DB accessed")):
            self.assertEqual(self.chatbot.query_excel_sheet("How many rows?", self.frame).iloc[0, 0], 3)
            pd.testing.assert_frame_equal(self.chatbot.query_excel_sheet("preview", self.frame), self.frame)
        self.client.chat.completions.create.assert_not_called()
        self.assertEqual(len(self.chatbot.query_excel_sheet("preview", pd.concat([self.frame] * 40))), 100)

    def test_arbitrary_columns_grouped_aggregates_and_bound_filter(self):
        result = self.ask("Total revenue per region ordered highest first", plan(
            columns=["c0"], group_by=["c0"], metrics=[{"function": "sum", "column": "c1"}],
            order_by=[{"column": "sum_c1", "direction": "desc"}],
        ))
        self.assertEqual(result.to_dict("list"), {"Region": ["West", "East"], "sum Revenue": [40, 20]})
        result = self.ask("West revenue above 10", plan(
            columns=["c1"], metrics=[], filter={"all": [
                {"column": "c0", "op": "eq", "value": "West"},
                {"column": "c1", "op": "gt", "value": 10},
            ]},
        ))
        self.assertEqual(result["Revenue"].tolist(), [30])
        prompt = self.client.chat.completions.create.call_args.kwargs["messages"][0]["content"]
        self.assertIn("Revenue", prompt)
        self.assertNotIn("West", prompt)

    def test_single_sheet_sql_trace_has_plan_statement_and_redacted_parameters(self):
        events = []
        token = self.chatbot.SQL_TRACE.set(self.chatbot.BackendTrace(events))
        self.addCleanup(self.chatbot.SQL_TRACE.reset, token)
        result = self.ask("Filtered revenue", plan(
            columns=["c1"], metrics=[], filter={"column": "c0", "op": "eq", "value": "West"},
        ))
        self.assertEqual(result["Revenue"].tolist(), [10, 30])
        self.assertEqual([event["event"] for event in events], [
            "planner_requested", "planner_completed", "source_loading", "query_started", "query_completed",
        ])
        self.assertIn("SELECT", events[-2]["sql"])
        self.assertEqual(events[-2]["route"], "sqlite_spreadsheet")
        self.assertTrue(events[-2]["parameters"])
        self.assertGreaterEqual(events[-1]["duration_ms"], 0)
        self.assertNotIn("West", self.chatbot.diagnostic_log(events))

    def test_dates_nulls_top_per_group_and_hostile_headers(self):
        result = self.ask("Missing dates", plan(filter={"column": "c2", "op": "is_null", "value": None}))
        self.assertEqual(result.iloc[0, 0], 1)
        result = self.ask("After January", plan(
            filter={"column": "c2", "op": "gte", "value": "2025-02-01"},
        ))
        self.assertEqual(result.iloc[0, 0], 1)
        result = self.ask("Highest revenue per region", plan(
            columns=["c0", "c1"], metrics=[], per_group=["c0"], limit_per_group=1,
            order_by=[{"column": "c1", "direction": "desc"}],
        ))
        self.assertEqual(result["Revenue"].tolist(), [30, 20])
        self.frame.columns = ["x]; DROP TABLE spreadsheet;--", "Revenue", "Date"]
        result = self.ask("Show first column", plan(columns=["c0"], metrics=[]))
        self.assertEqual(result.iloc[:, 0].tolist(), ["West", "East", "West"])

    def workbook_ask(self, question, query_plan, content=None):
        self.client.chat.completions.create.return_value = SimpleNamespace(choices=[
            SimpleNamespace(finish_reason="stop", message=SimpleNamespace(content=json.dumps(query_plan)))
        ])
        return self.chatbot.query_excel_workbook(question, self.frame, content or workbook_bytes(), "Sales")

    def test_workbook_planner_cross_sheet_lookup_schema_only_and_plain_answer(self):
        query = {"understood": True, "unhandled_conditions": [], "sheet": "Sales", "steps": [
            {"op": "lookup", "sheet": "Inventory", "left_on": ["Region"], "right_on": ["Product"],
             "columns": [{"column": "Stock", "name": "Available Stock"}]},
            {"op": "aggregate", "metrics": [{"name": "Median Revenue", "function": "median", "column": "Revenue"}]},
        ]}
        with patch.object(self.chatbot, "get_engine", side_effect=AssertionError("Employee DB accessed")):
            result = self.workbook_ask("What is the median revenue after looking up inventory?", query)
        self.assertEqual(result.iloc[0, 0], 20)
        self.assertEqual(result.attrs["source_sheets"], ["Sales", "Inventory"])
        self.assertIn("20", spreadsheet_answer(result))
        prompt = self.client.chat.completions.create.call_args.kwargs["messages"][0]["content"]
        self.assertIn("Inventory", prompt)
        self.assertIn("Stock", prompt)
        self.assertIn("median", prompt)
        self.assertNotIn("Widget", prompt)
        self.assertNotIn("West", prompt)
        self.assertNotIn("2025-01-01", prompt)

    def test_workbook_display_cap_is_after_calculation_and_unrelated_invalid_sheet_is_reported(self):
        content = workbook_bytes({
            "Sales": [("Region", "Revenue"), *[("North", 1) for _ in range(150)]],
            "Bad": [("Duplicate", " duplicate "), (1, 2)],
        })
        self.frame = read_excel_sheet(content, "Sales")
        result = self.workbook_ask("Sum the Revenue column", {
            "understood": True, "unhandled_conditions": [], "sheet": "Sales",
            "steps": [{"op": "aggregate", "metrics": [{"name": "Total", "function": "sum", "column": "Revenue"}]}],
        }, content)
        self.assertEqual(result.iloc[0, 0], 150)
        prompt = self.client.chat.completions.create.call_args.kwargs["messages"][0]["content"]
        self.assertIn('Unavailable sheets: {"Bad":', prompt)
        result = self.workbook_ask("List Revenue", {
            "understood": True, "unhandled_conditions": [], "sheet": "Sales",
            "steps": [{"op": "select", "columns": ["Revenue"]}],
        }, content)
        self.assertEqual(len(result), 150)

    def test_workbook_local_results_are_complete_while_legacy_preview_stays_capped(self):
        frame = pd.DataFrame({"Region": [f"Region {index}" for index in range(151)], "Sales": [1] * 151})
        for question in ("show all data", "total sales by region"):
            with self.subTest(question=question):
                result = self.chatbot.query_excel_workbook(question, frame, workbook_bytes(), "Sales")
                self.assertEqual(len(result), 151)
                self.assertEqual(len(self.chatbot.query_excel_sheet(question, frame)), 100)
        self.client.chat.completions.create.assert_not_called()

    def test_ui_full_excel_download_keeps_chat_and_table_at_100_rows(self):
        upload = BytesIO(workbook_bytes({"products": [
            ("Product ID", "Product Name"), *[(index, f"Product {index}") for index in range(1, 152)],
        ]}))
        upload.name = "complete.xlsx"
        with patch("streamlit.file_uploader", return_value=upload):
            app = employee_app().run(timeout=30)
            app.radio(key="data_source").set_value("Excel workbook").run(timeout=30)
            app.chat_input[0].set_value("show all data").run(timeout=30)
            self.assertEqual(len(app.exception), 0)
            message = app.session_state["chat_history"][0]
            self.assertEqual(len(message["result"]), 151)
            self.assertEqual(sorted(len(item.value) for item in app.dataframe), [10, 100])
            exported = pd.read_excel(BytesIO(message["excel_download"]))
            self.assertEqual(len(exported), 151)
            self.assertEqual(exported["Product ID"].iloc[-1], 151)
            self.assertIn("first 100 of 151", message["answer"])
            self.assertNotIn("Product Name: Product 101", message["answer"])
            app.run(timeout=30)
            self.assertEqual(app.session_state["chat_history"][0]["excel_download"], message["excel_download"])

    def test_ui_chat_preserves_actual_worksheet_loading_error_and_recovers_on_valid_sheet(self):
        upload = BytesIO(workbook_bytes({
            "products": [("Product ID", " product id "), (1, 2)],
            "Valid": [("Product ID", "Product Name"), (1, "Widget")],
        }))
        upload.name = "headers.xlsx"
        with patch("streamlit.file_uploader", return_value=upload):
            app = employee_app().run(timeout=30)
            app.radio(key="data_source").set_value("Excel workbook").run(timeout=30)
            app.chat_input[0].set_value("count rows").run(timeout=30)
            self.assertEqual(len(app.exception), 0)
            message = app.session_state["chat_history"][0]
            self.assertIn("headers must be unique", message["error"])
            self.assertIn("A1", message["error"])
            self.assertIn("B1", message["error"])
            self.assertEqual(message["excel_source"], "headers.xlsx / products")
            self.assertFalse(any("Results are limited to 100 rows" in item.value for item in app.caption))
            app.selectbox[0].select("Valid").run(timeout=30)
            app.chat_input[0].set_value("count rows").run(timeout=30)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(app.session_state["chat_history"][1]["result"].iloc[0, 0], 1)

    def test_workbook_plan_errors_missing_configuration_and_bounds(self):
        valid = {"understood": True, "unhandled_conditions": [], "sheet": "Sales",
                 "steps": [{"op": "select", "columns": ["Revenue"]}]}
        with patch.object(self.chatbot, "get_azure_openai_client", return_value=None):
            with self.assertRaisesRegex(ValueError, "Azure OpenAI"):
                self.chatbot.query_excel_workbook("Median Revenue", self.frame, workbook_bytes(), "Sales")
            self.assertEqual(self.chatbot.query_excel_workbook("count rows", self.frame, workbook_bytes(), "Sales").iloc[0, 0], 3)
        for invalid in ({}, {**valid, "unhandled_conditions": ["unsupported"]}, {**valid, "sheet": "Missing"}):
            with self.subTest(invalid=invalid), self.assertRaises(ValueError):
                self.workbook_ask("Unsupported workbook question", invalid)
        for content, finish in (("not JSON", "stop"), ("{}", "length")):
            self.client.chat.completions.create.return_value = SimpleNamespace(choices=[
                SimpleNamespace(finish_reason=finish, message=SimpleNamespace(content=content))
            ])
            with self.assertRaises(ValueError):
                self.chatbot.query_excel_workbook("Median Revenue", self.frame, workbook_bytes(), "Sales")
        with patch.object(self.chatbot, "MAX_SHEETS", 1):
            with self.assertRaisesRegex(ValueError, "20 worksheets"):
                self.workbook_ask("Median Revenue", valid)

    def test_exact_monthly_question_selects_sales_with_wrong_or_invalid_preview_without_ai(self):
        content = workbook_bytes({
            "products": [("Product ID", " product id "), (1, 2)],
            "Sales": [("Region", "Sales", "Date"),
                      ("PRIVATE_REGION", 100, "2025-01-15"), ("South", 200, "2025-02-10"),
                      ("North", 50, "2025-02-12")],
            "staff": [("Name", "Joining Date"), ("PRIVATE_PERSON", "2020-01-01")],
        })
        for selected, frame in (("products", None), ("staff", read_excel_sheet(content, "staff"))):
            with self.subTest(selected=selected), patch.object(self.chatbot, "get_azure_openai_client", return_value=None):
                result = self.chatbot.query_excel_workbook(
                    "Total Revenue by year and month of Order Date.", frame, content, selected,
                )
                self.assertEqual(result.to_dict("list"), {"Year": [2025, 2025], "Month": [1, 2], "Total Sales": [100, 250]})
                self.assertEqual(result.attrs["source_sheets"], ["Sales"])
                self.assertIn("'Date'", result.attrs["interpretation"])
        self.client.chat.completions.create.assert_not_called()

    def test_monthly_summary_rejects_ambiguity_and_does_not_discard_conditions(self):
        schemas = {
            "Sales": {"Revenue": "number", "Order Date": "date"},
            "Archive": {"Revenue": "number", "Order Date": "date"},
        }
        with self.assertRaisesRegex(ValueError, "More than one worksheet"):
            self.chatbot.monthly_workbook_plan("Total Revenue by year and month of Order Date.", schemas)
        plan, _ = self.chatbot.monthly_workbook_plan("Total Revenue by year and month of Order Date in Archive", schemas)
        self.assertEqual(plan["sheet"], "Archive")
        self.assertIsNone(self.chatbot.monthly_workbook_plan(
            "Total Revenue by year and month of Order Date where Region is North", schemas,
        ))
        self.assertIsNone(self.chatbot.monthly_workbook_plan(
            "Total Revenue by year and month of Order Date", {"staff": {"Revenue": "number", "Joining Date": "date"}},
        ))

    def test_named_row_count_and_sales_totals_route_away_from_preview(self):
        content = workbook_bytes()
        frame = read_excel_sheet(content, "Inventory")
        with patch.object(self.chatbot, "get_azure_openai_client", return_value=None):
            result = self.chatbot.query_excel_workbook("count rows in worksheet Sales", frame, content, "Inventory")
            self.assertEqual(result.iloc[0, 0], 3)
            self.assertEqual(result.attrs["source_sheets"], ["Sales"])
            result = self.chatbot.query_excel_workbook("total revenue by region", frame, content, "Inventory")
            self.assertEqual(result["Total Sales"].tolist(), [40, 20])
            self.assertEqual(result.attrs["source_sheets"], ["Sales"])

    def test_planner_rejection_explains_reason_and_available_columns(self):
        with self.assertRaises(ValueError) as caught:
            self.workbook_ask("Total Revenue by year and month of Shipment Date", {
                "understood": False, "unhandled_conditions": ["Shipment Date is missing; use Date instead."],
                "sheet": "Sales", "steps": [],
            })
        message = str(caught.exception)
        self.assertIn("Shipment Date is missing", message)
        self.assertIn("Sales: Region, Revenue, Date", message)
        self.assertIn("Inventory: Product, Stock", message)

    def test_planner_source_names_are_resolved_case_insensitively(self):
        result = self.workbook_ask("Median stock in inventory", {
            "understood": True, "unhandled_conditions": [], "sheet": "inventory", "steps": [
                {"op": "aggregate", "metrics": [{"name": "Median Stock", "function": "median", "column": "Stock"}]},
            ],
        })
        self.assertEqual(result.iloc[0, 0], 7)
        self.assertEqual(result.attrs["source_sheets"], ["Inventory"])

    def test_invalid_planner_response_is_diagnosed_without_logging_response_content(self):
        self.client.chat.completions.create.return_value = SimpleNamespace(choices=[
            SimpleNamespace(finish_reason="stop", message=SimpleNamespace(content="PRIVATE_MODEL_RESPONSE"))
        ])
        events = []
        with self.assertRaises(ValueError):
            self.chatbot.query_excel_workbook(
                "Median Revenue", self.frame, workbook_bytes(), "Sales", trace=self.chatbot.BackendTrace(events),
            )
        self.assertIn("planner_invalid_json", [entry["event"] for entry in events])
        self.assertNotIn("PRIVATE_MODEL_RESPONSE", self.chatbot.diagnostic_log(events))

    def test_ui_automatic_routing_from_invalid_preview_and_session_backend_log(self):
        upload = BytesIO(workbook_bytes({
            "products": [("ID", " id "), (1, 2)],
            "Sales": [("Revenue", "Order Date", "Region"),
                      (100, "2025-01-01", "PRIVATE_REGION"), (200, "2025-02-01", "PRIVATE_REGION")],
        }))
        upload.name = "private-workbook.xlsx"
        with patch("streamlit.file_uploader", return_value=upload):
            app = employee_app().run(timeout=30)
            app.radio(key="data_source").set_value("Excel workbook").run(timeout=30)
            app.selectbox[0].select("products").run(timeout=30)
            app.chat_input[0].set_value("Total Revenue by year and month of Order Date.").run(timeout=30)
            self.assertEqual(len(app.exception), 0)
            message = app.session_state["chat_history"][0]
            self.assertEqual(message["result"]["Total Revenue"].tolist(), [100, 200])
            self.assertEqual(message["excel_source"], "private-workbook.xlsx / Sales")
            events = message["backend_log"]
            self.assertIn("sheet_unavailable", [entry["event"] for entry in events])
            self.assertEqual(events[-1]["event"], "request_completed")
            text = self.chatbot.diagnostic_log(events)
            self.assertNotIn("PRIVATE_REGION", text)
            self.assertNotIn("Total Revenue by year and month", text)
            self.assertNotIn(upload.name, text)
            self.assertTrue(any(button.label == "Download backend log" for button in app.get("download_button")))
            count = len(events)
            app.run(timeout=30)
            self.assertEqual(len(app.session_state["chat_history"][0]["backend_log"]), count)
            app.button(key="new_chat").click().run(timeout=30)
            self.assertFalse(any(button.label == "Download backend log" for button in app.get("download_button")))

    def test_workbook_advanced_ui_caches_answer_and_discards_input_snapshots(self):
        self.chatbot.st.cache_resource.clear()
        self.addCleanup(self.chatbot.st.cache_resource.clear)
        upload = BytesIO(workbook_bytes())
        upload.name = "advanced.xlsx"
        self.client.chat.completions.create.return_value = SimpleNamespace(choices=[
            SimpleNamespace(finish_reason="stop", message=SimpleNamespace(content=json.dumps({
                "understood": True, "unhandled_conditions": [], "sheet": "Inventory", "steps": [
                    {"op": "aggregate", "metrics": [{"name": "Median Stock", "function": "median", "column": "Stock"}]},
                ],
            })))
        ])
        with patch("streamlit.file_uploader", return_value=upload), patch(
            "openai.AzureOpenAI", return_value=self.client,
        ), patch.dict(os.environ, {
            "AZURE_OPENAI_ENDPOINT": "https://example.test", "AZURE_OPENAI_API_KEY": "test",
            "AZURE_OPENAI_API_VERSION": "2024-10-21",
        }), patch.object(self.chatbot, "get_engine", side_effect=AssertionError("Employee DB accessed")):
            app = employee_app().run(timeout=30)
            app.radio(key="data_source").set_value("Excel workbook").run(timeout=30)
            app.chat_input[0].set_value("What is the median Stock in Inventory?").run(timeout=30)
            self.assertEqual(len(app.exception), 0)
            self.client.chat.completions.create.assert_called_once()
            message = app.session_state["chat_history"][0]
            self.assertIn("result", message, message.get("error"))
            self.assertEqual(message["result"].iloc[0, 0], 7)
            self.assertEqual(message["excel_source"], "advanced.xlsx / Inventory")
            self.assertNotIn("excel_frame", message)
            self.assertNotIn("excel_content", message)
            self.assertNotIn("excel_sheet", message)
            self.assertIn("7", message["answer"])
            self.client.chat.completions.create.reset_mock()
            app.radio(key="data_source").set_value("Employee database").run(timeout=30)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(app.session_state["chat_history"][0]["result"].iloc[0, 0], 7)
            self.client.chat.completions.create.assert_not_called()
            self.assertEqual(sum(button.label == "Download Excel" for button in app.get("download_button")), 1)

    def test_query_limits_and_unsupported_plans(self):
        self.frame = pd.concat([self.frame] * 40)
        self.assertEqual(len(self.ask("List revenue", plan(columns=["c1"], metrics=[]))), 100)
        for invalid in (
            plan(columns=["unknown"]), plan(understood=False), plan(unhandled_conditions=["join sheets"]),
            plan(limit=101), plan(metrics=[{"function": "sum", "column": "c0"}]),
            plan(filter={"column": "c1", "op": "gt", "value": "10"}),
        ):
            with self.subTest(plan=invalid), self.assertRaisesRegex(ValueError, "couldn't fully interpret"):
                self.ask("Unsupported", invalid)

    def test_missing_azure_invalid_json_and_incomplete_responses(self):
        with patch.object(self.chatbot, "get_azure_openai_client", return_value=None):
            with self.assertRaisesRegex(ValueError, "Azure OpenAI"):
                self.chatbot.query_excel_sheet("Average revenue", self.frame)
        for content, finish in (("not JSON", "stop"), ("{}", "length")):
            self.client.chat.completions.create.return_value = SimpleNamespace(choices=[
                SimpleNamespace(finish_reason=finish, message=SimpleNamespace(content=content))
            ])
            with self.subTest(content=content), self.assertRaises(ValueError):
                self.chatbot.query_excel_sheet("Average revenue", self.frame)

    def test_natural_language_north_sales_total_without_azure(self):
        sales = pd.DataFrame({
            "Region": ["North", " north ", "South", "North"],
            "Sales": [1250.5, 250.25, 9000, None],
        })
        with patch.object(self.chatbot, "get_azure_openai_client", return_value=None):
            result = self.chatbot.query_excel_sheet("What are the total sales for the North region?", sales)
        self.assertEqual(result.iloc[0, 0], 1500.75)
        self.client.chat.completions.create.assert_not_called()

    def test_grouped_sales_question_runs_locally_and_matches_written_summary(self):
        with patch.object(self.chatbot, "get_azure_openai_client", return_value=None):
            for question in (
                "What are the total sales by region?", "Total revenue per region",
                "Show total sales for each region",
            ):
                with self.subTest(question=question):
                    result = self.chatbot.query_excel_sheet(question, self.frame)
                    self.assertEqual(result.to_dict("list"), {"Region": ["West", "East"], "Total Sales": [40, 20]})
        _, updated = write_region_summary(workbook_bytes(), confirmed=True)
        summary = pd.read_excel(BytesIO(updated), sheet_name="Region_Summary")
        pd.testing.assert_frame_equal(result, summary, check_dtype=False)
        self.client.chat.completions.create.assert_not_called()
        many_regions = pd.DataFrame({"Region": [f"Region {i}" for i in range(101)], "Sales": [1] * 101})
        self.assertEqual(len(self.chatbot.query_excel_sheet("Total sales by region", many_regions)), 100)

    def test_sales_ui_answers_in_plain_language_and_preserves_download(self):
        upload = BytesIO(workbook_bytes({
            "Inventory": [("Product", "Stock"), ("Widget", 7)],
            "Sales": [("Region", "Sales"), ("North", 1250.5), ("South", 9000), ("North", 250.25)],
        }))
        upload.name = "sales.xlsx"
        with patch("streamlit.file_uploader", return_value=upload), patch(
            "socket.create_connection", side_effect=AssertionError("Unexpected network"),
        ):
            app = employee_app().run(timeout=30)
            app.radio(key="data_source").set_value("Excel workbook").run(timeout=30)
            self.assertEqual(app.selectbox[0].value, "Sales")
            self.assertFalse(any("Chat commands:" in caption.value or "Choose one worksheet" in caption.value
                                 for caption in app.sidebar.caption))
            app.chat_input[0].set_value("What are the total sales for the North region?").run(timeout=30)
            self.assertEqual(len(app.exception), 0)
            message = app.session_state["chat_history"][-1]
            self.assertEqual(message["answer"], "The total sales for the North region is 1,500.75.")
            self.assertTrue(any(message["answer"] == item.value for item in app.markdown))
            self.assertEqual(pd.read_excel(BytesIO(message["excel_download"])).iloc[0, 0], 1500.75)
            route = next(event for event in message["backend_log"] if event["event"] == "route_selected")
            functions = {item["function"] for item in route["implementation"]}
            self.assertEqual(functions, {"local_excel_result", "query_sales_total"})
            self.assertTrue(all(caption.value.startswith("Sources used:") for caption in app.caption))
            self.assertEqual(len(app.caption), 1)
            self.assertIn("EXCEL:", app.caption[0].value)
            implementation = {
                item["function"]: item["code"]
                for event in message["backend_log"] for item in event.get("implementation", [])
            }
            for function in ("query_sales_total", "spreadsheet_answer", "generate_excel_download"):
                self.assertIn(function, implementation)
                self.assertIn(f"def {function}(", implementation[function])
            app.run(timeout=30)
            self.assertEqual(app.session_state["chat_history"][-1]["answer"], message["answer"])

    def test_ui_sheet_selection_snapshots_and_database_switch(self):
        upload = BytesIO(workbook_bytes())
        upload.name = "sample.xlsx"
        with patch("streamlit.file_uploader", return_value=upload), patch.object(
            self.chatbot, "get_engine", side_effect=AssertionError("Employee DB accessed"),
        ):
            app = employee_app().run(timeout=30)
            app.radio(key="data_source").set_value("Excel workbook").run(timeout=30)
            app.chat_input[0].set_value("count rows").run(timeout=30)
            self.assertEqual(len(app.exception), 0)
            first = app.session_state["chat_history"][0]
            self.assertEqual(first["result"].iloc[0, 0], 3)
            self.assertNotIn("excel_frame", first)
            app.selectbox[0].select("Inventory").run(timeout=30)
            app.chat_input[0].set_value("preview").run(timeout=30)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(app.session_state["chat_history"][1]["result"]["Stock"].tolist(), [7])
            self.assertIn("Sales", app.session_state["chat_history"][0]["excel_source"])
            app.radio(key="data_source").set_value("Employee database").run(timeout=30)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(sum(button.label == "Download Excel" for button in app.get("download_button")), 2)
            app.button(key="new_chat").click().run(timeout=30)
            self.assertEqual(len(app.session_state["chat_history"]), 0)

    def test_ui_missing_and_invalid_workbooks_never_fall_back_to_employees(self):
        with patch("streamlit.file_uploader", return_value=None), patch.object(
            self.chatbot, "get_engine", side_effect=AssertionError("Employee DB accessed"),
        ):
            app = employee_app().run(timeout=30)
            app.radio(key="data_source").set_value("Excel workbook").run(timeout=30)
            app.chat_input[0].set_value("count rows").run(timeout=30)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.session_state["chat_history"]), 1)
            self.assertIn("Upload a valid Excel", app.session_state["chat_history"][0]["error"])
            self.assertTrue(any("Upload a valid Excel" in warning.value for warning in app.warning))
        invalid = BytesIO(b"invalid")
        invalid.name = "invalid.xlsx"
        with patch("streamlit.file_uploader", return_value=invalid):
            app.run(timeout=30)
            self.assertEqual(len(app.exception), 0)
            self.assertTrue(any("valid, unencrypted" in warning.value for warning in app.warning))

    def test_excel_results_are_not_sent_as_general_chat_context(self):
        context = self.chatbot.conversation_context([
            {"kind": "excel", "question": "preview", "answer": "Private cell value"},
        ])
        self.assertIn(("user", "preview"), context)
        self.assertNotIn("Private cell value", str(context))

    def test_region_summary_ui_requires_confirmation_and_offers_both_downloads(self):
        upload = BytesIO(workbook_bytes())
        upload.name = "sample.xlsx"
        with patch("streamlit.file_uploader", return_value=upload):
            app = employee_app().run(timeout=30)
            app.radio(key="data_source").set_value("Excel workbook").run(timeout=30)
            create = next(button for button in app.button if button.label == "Create Region_Summary")
            self.assertTrue(create.disabled)
            cache = app.session_state["excel_workbook"]
            self.assertEqual(cache["backup"], upload.getvalue())
            self.assertNotIn("region_summary_output", cache)
            self.assertTrue(any(button.label == "Download original backup" for button in app.get("download_button")))
            app.checkbox[0].check().run(timeout=30)
            app.button(key=create.key).click().run(timeout=30)
            self.assertEqual(len(app.exception), 0)
            cache = app.session_state["excel_workbook"]
            output = cache["region_summary_output"]
            self.assertEqual(load_workbook(BytesIO(output)).sheetnames, ["Sales", "Inventory", "Region_Summary"])
            self.assertEqual(cache["backup"], upload.getvalue())
            self.assertTrue(any(button.label == "Download workbook with Region_Summary"
                                for button in app.get("download_button")))
            app.run(timeout=30)
            self.assertEqual(app.session_state["excel_workbook"]["region_summary_output"], output)

    def test_region_summary_ui_refuses_to_replace_existing_sheet(self):
        upload = BytesIO(workbook_bytes({
            "Sales": [("Region", "Sales"), ("North", 10)],
            "Region_Summary": [("Original",), (99,)],
        }))
        upload.name = "existing.xlsx"
        with patch("streamlit.file_uploader", return_value=upload):
            app = employee_app().run(timeout=30)
            app.radio(key="data_source").set_value("Excel workbook").run(timeout=30)
            app.checkbox[0].check().run(timeout=30)
            create = next(button for button in app.button if button.label == "Create Region_Summary")
            create.click().run(timeout=30)
            self.assertEqual(len(app.exception), 0)
            self.assertTrue(any("already exists" in error.value for error in app.error))
            self.assertNotIn("region_summary_output", app.session_state["excel_workbook"])

    def test_missing_sales_and_access_errors_display_friendly_ui_messages(self):
        upload = BytesIO(workbook_bytes({"Inventory": [("Product",), ("Widget",)]}))
        upload.name = "inventory.xlsx"
        with patch("streamlit.file_uploader", return_value=upload):
            app = employee_app().run(timeout=30)
            app.radio(key="data_source").set_value("Excel workbook").run(timeout=30)
            self.assertEqual(len(app.exception), 0)
            self.assertTrue(any("doesn't have a Sales sheet" in info.value for info in app.info))
        locked = BytesIO(workbook_bytes())
        locked.name = "locked.xlsx"
        with patch("streamlit.file_uploader", return_value=locked), patch(
            "excel_support.load_workbook", side_effect=PermissionError("PRIVATE_TECHNICAL_ERROR"),
        ):
            app.run(timeout=30)
            self.assertEqual(len(app.exception), 0)
            self.assertTrue(any("may be locked" in warning.value for warning in app.warning))
            self.assertFalse(any("PRIVATE_TECHNICAL_ERROR" in warning.value for warning in app.warning))

    def test_workbook_commands_open_and_close_controls_without_a_file(self):
        with patch("streamlit.file_uploader", return_value=None):
            app = employee_app().run(timeout=30)
            app.chat_input[0].set_value("open Excel").run(timeout=30)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(app.radio(key="data_source").value, "Excel workbook")
            self.assertIn("Upload your workbook", app.session_state["chat_history"][-1]["answer"])
            app.chat_input[0].set_value("open Sales").run(timeout=30)
            self.assertEqual(len(app.exception), 0)
            self.assertIn("Upload an Excel workbook first", app.session_state["chat_history"][-1]["answer"])
            app.chat_input[0].set_value("close Excel").run(timeout=30)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(app.radio(key="data_source").value, "Employee database")
            self.assertNotIn("excel_workbook", app.session_state)
            self.assertEqual(app.session_state["excel_uploader_version"], 1)

    def test_prompt_sheet_selection_and_recall_survive_close_and_rerun(self):
        upload = BytesIO(workbook_bytes())
        upload.name = "sample.xlsx"
        with patch("streamlit.file_uploader", return_value=upload):
            app = employee_app().run(timeout=30)
            app.chat_input[0].set_value("open Excel").run(timeout=30)
            app.chat_input[0].set_value("open Inventory").run(timeout=30)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(app.selectbox[0].value, "Inventory")
            self.assertEqual(app.selectbox[0].proto.default, 1)
            app.chat_input[0].set_value("count rows").run(timeout=30)
            self.assertEqual(app.session_state["chat_history"][-1]["result"].iloc[0, 0], 1)
            app.chat_input[0].set_value("what's the last question?").run(timeout=30)
            self.assertEqual(app.session_state["chat_history"][-1]["answer"],
                             'Your previous question in this chat was: "count rows"')
            app.chat_input[0].set_value("whats the last question?").run(timeout=30)
            self.assertIn('"count rows"', app.session_state["chat_history"][-1]["answer"])
            app.run(timeout=30)
            self.assertEqual(len(app.exception), 0)
            app.chat_input[0].set_value("close Excel").run(timeout=30)
            self.assertNotIn("excel_workbook", app.session_state)
            app.chat_input[0].set_value("what was my last question?").run(timeout=30)
            self.assertIn('"count rows"', app.session_state["chat_history"][-1]["answer"])
            self.assertEqual(len(app.exception), 0)
        with patch("streamlit.file_uploader", return_value=None):
            app.chat_input[0].set_value("open Excel").run(timeout=30)
            self.assertNotIn("excel_workbook", app.session_state)
            self.assertIn("Upload your workbook", app.session_state["chat_history"][-1]["answer"])

    def test_open_sheet_restores_remembered_workbook_after_switching_sources(self):
        upload = BytesIO(workbook_bytes())
        upload.name = "remembered.xlsx"
        with patch("streamlit.file_uploader", return_value=upload):
            app = employee_app().run(timeout=30)
            app.radio(key="data_source").set_value("Excel workbook").run(timeout=30)
            app.radio(key="data_source").set_value("Employee database").run(timeout=30)
        with patch("streamlit.file_uploader", return_value=None):
            app.chat_input[0].set_value("open Inventory").run(timeout=30)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(app.radio(key="data_source").value, "Excel workbook")
            self.assertEqual(app.selectbox[0].value, "Inventory")
            self.assertEqual(app.session_state["excel_workbook"]["file_name"], "remembered.xlsx")
            app.chat_input[0].set_value("open Missing").run(timeout=30)
            self.assertIn("Available worksheets: Sales, Inventory", app.session_state["chat_history"][-1]["answer"])
            self.assertEqual(app.selectbox[0].value, "Inventory")

    def test_last_question_works_with_invalid_headers_and_without_a_workbook(self):
        with patch("streamlit.file_uploader", return_value=None):
            app = employee_app().run(timeout=30)
            app.radio(key="data_source").set_value("Excel workbook").run(timeout=30)
            app.chat_input[0].set_value("whats the last question?").run(timeout=30)
            self.assertEqual(len(app.exception), 0)
            self.assertIn("no earlier messages", app.session_state["chat_history"][-1]["answer"])
        upload = BytesIO(workbook_bytes({"Sales": [("Region", "region"), ("North", 100)]}))
        upload.name = "invalid-headers.xlsx"
        with patch("streamlit.file_uploader", return_value=upload):
            app.run(timeout=30)
            self.assertTrue(any("headers must be unique" in warning.value for warning in app.warning))
            question = "What are the total sales by region?"
            app.chat_input[0].set_value(question).run(timeout=30)
            self.assertIn("error", app.session_state["chat_history"][-1])
            app.chat_input[0].set_value("what is the last question?").run(timeout=30)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(app.session_state["chat_history"][-1]["kind"], "recall")
            self.assertIn(question, app.session_state["chat_history"][-1]["answer"])
            app.button(key="new_chat").click().run(timeout=30)
            app.chat_input[0].set_value("what's the last question?").run(timeout=30)
            self.assertIn("no earlier messages", app.session_state["chat_history"][-1]["answer"])

    def test_command_and_recall_parsers_do_not_need_azure(self):
        self.assertEqual(self.chatbot.parse_workbook_command("please open Excel"), ("open", None))
        self.assertEqual(self.chatbot.parse_workbook_command("close the workbook"), ("close", None))
        self.assertEqual(self.chatbot.parse_workbook_command('open sheet "Sales"'), ("sheet", "Sales"))
        self.assertIsNone(self.chatbot.parse_workbook_command("How many employees do we have?"))
        for question in ("what's the last question?", "whats the last question?", "what is the last question?"):
            with self.subTest(question=question):
                self.assertTrue(self.chatbot.is_previous_question_request(question))
        self.assertIn('"count rows"', self.chatbot.recall_previous_question([
            {"question": "count rows"}, {"kind": "workbook_command", "question": "close Excel"},
            {"kind": "recall", "question": "whats the last question?"},
        ]))

    def test_open_products_worksheet_accepts_suffix_and_prefix_forms(self):
        upload = BytesIO(workbook_bytes({
            "products": [("Product", "Stock"), ("Widget", 7)],
            "Sales": [("Region", "Sales"), ("North", 100)],
            "staff": [("Name", "Role"), ("Asha", "Analyst")],
        }))
        upload.name = "command-forms.xlsx"
        with patch("streamlit.file_uploader", return_value=upload):
            app = employee_app().run(timeout=30)
            app.chat_input[0].set_value("open Excel").run(timeout=30)
            for question, expected in (
                ("open worksheet staff", "staff"),
                ("open products worksheet", "products"),
                ("open Sales sheet", "Sales"),
                ('please open the "products" worksheet', "products"),
                ("switch to STAFF worksheet", "staff"),
            ):
                with self.subTest(question=question):
                    app.chat_input[0].set_value(question).run(timeout=30)
                    self.assertEqual(len(app.exception), 0)
                    self.assertEqual(app.selectbox[0].value, expected)
                    self.assertIn(f"Selected worksheet {expected}", app.session_state["chat_history"][-1]["answer"])
            app.chat_input[0].set_value("open Missing worksheet").run(timeout=30)
            self.assertIn("Available worksheets: products, Sales, staff", app.session_state["chat_history"][-1]["answer"])
            self.assertEqual(app.selectbox[0].value, "staff")

    def test_sheet_resolution_prefers_literal_names_before_optional_suffixes(self):
        resolve = self.chatbot.resolve_workbook_sheet
        self.assertEqual(resolve("products worksheet", ["products", "products worksheet"]), "products worksheet")
        self.assertEqual(resolve("Products WORKSHEET", ["products"]), "products")
        self.assertEqual(resolve('products" worksheet', ["products"]), "products")
        self.assertEqual(resolve("Sales sheet", ["Sales"]), "Sales")
        self.assertIsNone(resolve("Missing worksheet", ["products", "Sales", "staff"]))


if __name__ == "__main__":
    unittest.main()
