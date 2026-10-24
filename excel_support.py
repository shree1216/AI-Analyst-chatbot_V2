from collections import Counter
from datetime import date, datetime
from io import BytesIO
import math
from numbers import Real
from pathlib import Path
import re
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils.cell import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet
import pandas as pd


EXCEL_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_EXCEL_BYTES = 10_000_000
MAX_EXCEL_ROWS = 50_000
MAX_EXCEL_COLUMNS = 100
SALES_SHEET_MISSING = (
    "I couldn't find a sheet named 'Sales' in this workbook. "
    "Add or rename the sales worksheet to 'Sales', save the file, and upload it again."
)


def _excel_access_message(error: OSError) -> str:
    if isinstance(error, FileNotFoundError):
        return "I couldn't find the Excel file. Check the file path and try again."
    if isinstance(error, IsADirectoryError):
        return "That path points to a folder. Choose an .xlsx file instead."
    if isinstance(error, PermissionError) or getattr(error, "winerror", None) in {32, 33}:
        return (
            "I couldn't access the Excel file. It may be locked or you may not have permission. "
            "Close it in Excel or other apps, check your file permissions, and try again."
        )
    return "I couldn't access the Excel file. Check that it is available and try again."


def _excel_content(source: bytes | str | Path) -> bytes:
    if isinstance(source, bytes):
        content = source
    else:
        try:
            path = Path(source)
            if path.is_dir():
                raise IsADirectoryError()
            with path.open("rb") as workbook_file:
                content = workbook_file.read(MAX_EXCEL_BYTES + 1)
        except OSError as error:
            raise ValueError(_excel_access_message(error)) from error
        except ValueError as error:
            raise ValueError("That file path isn't valid. Choose an .xlsx file and try again.") from error
    if not content or len(content) > MAX_EXCEL_BYTES:
        raise ValueError("Upload an .xlsx workbook smaller than 10 MB.")
    return content


def open_excel_workbook(content: bytes | str | Path) -> Workbook:
    content = _excel_content(content)
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 50_000_000:
                raise ValueError("The expanded workbook must be smaller than 50 MB.")
        return load_workbook(BytesIO(content), read_only=True, data_only=True)
    except OSError as error:
        raise ValueError(_excel_access_message(error)) from error
    except (BadZipFile, InvalidFileException, ParseError, KeyError) as error:
        raise ValueError("Could not read this workbook. Upload a valid, unencrypted .xlsx file.") from error


def excel_sheet_names(content: bytes | str | Path) -> list[str]:
    workbook = open_excel_workbook(content)
    try:
        names = [worksheet.title for worksheet in workbook.worksheets]
        if not names:
            raise ValueError("The workbook has no data worksheets.")
        return names
    finally:
        workbook.close()


def read_excel_sheet(content: bytes | str | Path, sheet: str) -> pd.DataFrame:
    workbook = open_excel_workbook(content)
    try:
        if sheet not in [worksheet.title for worksheet in workbook.worksheets]:
            raise ValueError(
                SALES_SHEET_MISSING if sheet.casefold() == "sales"
                else "I couldn't find the selected worksheet. Choose another worksheet and try again."
            )
        worksheet = workbook[sheet]
        worksheet.reset_dimensions()
        rows = worksheet.iter_rows(values_only=True)
        header = list(next(rows, ()))
        while header and header[-1] is None:
            header.pop()
        if not header or len(header) > MAX_EXCEL_COLUMNS:
            raise ValueError(f"Worksheet '{sheet}': use a header row with 1 to 100 columns in row 1.")
        if any(not isinstance(value, str) or not value.strip() for value in header):
            cells = [f"{get_column_letter(index + 1)}1" for index, value in enumerate(header)
                     if not isinstance(value, str) or not value.strip()]
            raise ValueError(
                f"Worksheet '{sheet}': every column needs a nonempty text header in the first row. "
                f"Check {', '.join(cells)}. Put the actual column headings in row 1, save, and upload again."
            )
        columns = [value.strip() for value in header]
        counts = Counter(value.casefold() for value in columns)
        if len(counts) != len(columns):
            cells = [f"{get_column_letter(index + 1)}1" for index, value in enumerate(columns)
                     if counts[value.casefold()] > 1]
            raise ValueError(
                f"Worksheet '{sheet}': column headers must be unique (ignoring case and surrounding spaces). "
                f"Repeated headers are at {', '.join(cells)}. Give them distinct names, save, and upload again."
            )
        values = []
        for row_number, row in enumerate(rows, start=2):
            if len(row) > len(columns) and any(value is not None for value in row[len(columns):]):
                column = next(index for index, value in enumerate(row, start=1) if index > len(columns) and value is not None)
                raise ValueError(
                    f"Worksheet '{sheet}': some cells have no column header, for example "
                    f"{get_column_letter(column)}{row_number}. Add a distinct text header in row 1, save, and upload again."
                )
            if len(values) >= MAX_EXCEL_ROWS:
                raise ValueError("A worksheet can contain at most 50,000 data rows.")
            values.append(tuple(row[:len(columns)]) + (None,) * max(0, len(columns) - len(row)))
        frame = pd.DataFrame(values, columns=columns).dropna(how="all").reset_index(drop=True)
        if frame.empty:
            raise ValueError("The selected worksheet has no data rows.")
        return frame
    except OSError as error:
        raise ValueError(_excel_access_message(error)) from error
    finally:
        workbook.close()


def spreadsheet_schema(frame: pd.DataFrame) -> tuple[pd.DataFrame, dict[str, str], dict[str, str]]:
    data = frame.copy()
    names = {f"c{index}": str(column) for index, column in enumerate(frame.columns)}
    data.columns = list(names)
    schema = {}
    for column in data:
        series = data[column]
        populated = series.dropna()
        if pd.api.types.is_numeric_dtype(series) and not pd.api.types.is_bool_dtype(series):
            schema[column] = "number"
        elif not populated.empty and all(isinstance(value, (date, datetime)) for value in populated):
            schema[column] = "date"
            data[column] = series.map(lambda value: None if pd.isna(value) else value.date().isoformat()
                                      if isinstance(value, datetime) else value.isoformat())
        else:
            schema[column] = "text"
            data[column] = series.map(lambda value: None if pd.isna(value) else str(value))
    return data, schema, names


def _sales_column(frame: pd.DataFrame, metric: str = "sales") -> str | None:
    headers = {str(column): re.sub(r"[_\s]+", " ", str(column)).strip().casefold() for column in frame}
    candidates = [column for column, header in headers.items() if header == metric]
    if not candidates:
        candidates = [column for column, header in headers.items()
                      if header in {"sales", "total sales", "sales amount", "revenue"}]
    return candidates[0] if len(candidates) == 1 else None


def _sales_values(frame: pd.DataFrame, column: str) -> pd.Series:
    series = frame[column]
    if not series.dropna().empty and (
        not pd.api.types.is_numeric_dtype(series) or pd.api.types.is_bool_dtype(series)
    ):
        raise ValueError(f"The {column} column must contain numeric values to calculate total sales.")
    return series


def region_sales_totals(frame: pd.DataFrame, metric: str = "sales") -> pd.DataFrame:
    column = _sales_column(frame, metric)
    regions = [name for name in frame if str(name).strip().casefold() == "region"]
    if column is None or len(regions) != 1:
        raise ValueError("The Sales sheet needs a Region column and one unambiguous sales column.")
    values = _sales_values(frame, column)
    labels = frame[regions[0]].astype("string").str.strip()
    labels = labels.mask(labels.eq(""), pd.NA)
    keys = labels.str.casefold()
    totals = values.groupby(keys, dropna=False, sort=False).sum(min_count=1)
    display_labels = labels.groupby(keys, dropna=False, sort=False).first()
    return pd.DataFrame({"Region": display_labels.to_numpy(), "Total Sales": totals.to_numpy()})


def query_sales_total(question: str, frame: pd.DataFrame) -> pd.DataFrame | None:
    match = re.fullmatch(
        r"(?:(?:what (?:are|is)|show(?: me)?|give(?: me)?)\s+)?(?:the\s+)?"
        r"(?:total|sum(?: of)?)\s+(?P<metric>sales amount|total sales|sales|revenue)"
        r"(?:\s+(?:(?P<grouped>(?:by|per|for each)\s+region)|"
        r"(?:for|in)\s+(?:the\s+)?(?P<region>.+?)\s+region))?",
        question.strip().rstrip("?.!").strip(), flags=re.IGNORECASE,
    )
    if match is None:
        return None
    headers = {column: re.sub(r"[_\s]+", " ", str(column)).strip().casefold() for column in frame}
    metric = match["metric"].casefold()
    if match["grouped"]:
        return region_sales_totals(frame, metric)
    column = _sales_column(frame, metric)
    if column is None:
        return None
    selected = frame
    scope = ""
    if match["region"]:
        regions = [column for column, header in headers.items() if header == "region"]
        if len(regions) != 1:
            return None
        region_column = regions[0]
        region = match["region"].strip()
        values = frame[region_column].astype("string").str.strip()
        mask = values.str.casefold().eq(region.casefold()).fillna(False)
        if not mask.any():
            return None
        selected = frame.loc[mask]
        region_label = str(values.loc[mask].iloc[0])
        scope = f" for the {region_label} region"
    series = _sales_values(selected, column)
    result = pd.DataFrame({f"sum {column}": [series.sum(min_count=1)]})
    result.attrs["answer_subject"] = f"The total {metric.removeprefix('total ')}{scope}"
    return result


def write_region_summary(
    content: bytes | str | Path, *, confirmed: bool = False,
) -> tuple[bytes, bytes]:
    """Return an original backup and an updated copy; never modify the source."""
    if confirmed is not True:
        raise ValueError("Please confirm before writing Region_Summary.")
    backup = memoryview(_excel_content(content)).tobytes()
    names = excel_sheet_names(backup)
    if any(name.casefold() == "region_summary" for name in names):
        raise ValueError("Region_Summary already exists. No existing sheet will be overwritten.")
    source = next((name for name in names if name.casefold() == "sales"), None)
    if source is None:
        raise ValueError(SALES_SHEET_MISSING)
    frame = read_excel_sheet(backup, source)
    summary = region_sales_totals(frame)
    try:
        workbook = load_workbook(BytesIO(backup), data_only=False)
    except OSError as error:
        raise ValueError(_excel_access_message(error)) from error
    try:
        if any(name.casefold() == "region_summary" for name in workbook.sheetnames):
            raise ValueError("Region_Summary already exists. No existing sheet will be overwritten.")
        worksheet = workbook.create_sheet("Region_Summary")
        worksheet.append(["Region", "Total Sales"])
        for label, total in summary.itertuples(index=False, name=None):
            worksheet.append([
                None if pd.isna(label) else str(label),
                None if pd.isna(total) else total,
            ])
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        _literal_excel_text(worksheet)
        output = BytesIO()
        workbook.save(output)
        return backup, output.getvalue()
    except OSError as error:
        raise ValueError(_excel_access_message(error)) from error
    finally:
        workbook.close()


def _literal_excel_text(worksheet: Worksheet) -> None:
    for row in worksheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
                cell.hyperlink = None


def spreadsheet_answer(result: pd.DataFrame) -> str:
    def value_text(value: object) -> str:
        if pd.isna(value):
            return "not available"
        if isinstance(value, Real) and not isinstance(value, bool):
            if math.isfinite(value) and value == int(value):
                return f"{int(value):,}"
            return format(value, ",")
        if isinstance(value, (date, datetime)):
            return value.isoformat()
        return str(value)

    def label_text(column: object) -> str:
        label = str(column)
        for prefix, description in (
            ("sum ", "total "), ("average ", "average "), ("min ", "minimum "),
            ("max ", "maximum "), ("count distinct ", "distinct count of "), ("count ", "count of "),
        ):
            if label.startswith(prefix):
                return description + label[len(prefix):]
        return label.replace("_", " ")

    if result.empty:
        return "No matching rows were found in this worksheet."
    if result.shape == (1, 1):
        value = result.iloc[0, 0]
        if str(result.columns[0]) == "row_count":
            return f"There {'is' if value == 1 else 'are'} {value_text(value)} {'row' if value == 1 else 'rows'}."
        subject = result.attrs.get("answer_subject") or f"The {label_text(result.columns[0])}"
        if pd.isna(value):
            return f"No value is available for {subject.removeprefix('The ')}."
        return f"{subject} is {value_text(value)}."
    lines = [
        "; ".join(f"{label_text(column)}: {value_text(value)}" for column, value in zip(result.columns, row)) + "."
        for row in result.itertuples(index=False, name=None)
    ]
    if len(lines) == 1:
        return lines[0]
    return f"Here are the {len(result):,} matching results:\n\n" + "\n".join(f"- {line}" for line in lines)


def generate_excel_download(tables: list[tuple[str, pd.DataFrame]]) -> bytes:
    output = BytesIO()
    used_names = set()
    workbook = Workbook(write_only=True)
    try:
        for title, table in tables:
            base = re.sub(r"[\[\]:*?/\\]", "_", title).strip("'")[:31] or "Sheet"
            name = base
            suffix = 2
            while name.casefold() in used_names:
                tail = f" ({suffix})"
                name = base[:31 - len(tail)] + tail
                suffix += 1
            used_names.add(name.casefold())
            worksheet = workbook.create_sheet(name)
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = f"A1:{get_column_letter(max(1, len(table.columns)))}{len(table) + 1}"
            headers = []
            for column in table.columns:
                cell = WriteOnlyCell(worksheet, value=str(column))
                cell.data_type = "s"
                cell.font = Font(bold=True)
                headers.append(cell)
            worksheet.append(headers)
            for row in table.itertuples(index=False, name=None):
                cells = []
                for value in row:
                    if isinstance(value, (list, dict)):
                        value = str(value)
                    if pd.isna(value):
                        value = None
                    elif isinstance(value, datetime) and value.tzinfo is not None:
                        value = value.replace(tzinfo=None)
                    elif isinstance(value, Real) and not math.isfinite(value):
                        value = str(value)
                    cell = WriteOnlyCell(worksheet, value=value)
                    if isinstance(value, str):
                        cell.data_type = "s"
                    cells.append(cell)
                worksheet.append(cells)
        if not tables:
            raise ValueError("Choose at least one result table to download.")
        workbook.save(output)
    finally:
        workbook.close()
    return output.getvalue()
